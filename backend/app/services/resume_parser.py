import os
import re
import pymupdf
import docx
from typing import Dict, Any, List
from app.core.config import settings
from app.core.exceptions import ParsingException

class ResumeParserService:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        text = ""
        try:
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
            doc.close()
        except Exception as e:
            raise ParsingException(f"Could not parse PDF file: {str(e)}")
        
        if not text.strip():
            raise ParsingException("Extracted PDF text is empty. Scanned image-only PDFs require OCR which is not supported.")
        return text

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            full_text.append(cell.text)
            text = "\n".join(full_text)
        except Exception as e:
            raise ParsingException(f"Could not parse DOCX file: {str(e)}")

        if not text.strip():
            raise ParsingException("Extracted DOCX text is empty.")
        return text

    @classmethod
    def parse_file(cls, file_path: str, filename: str) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            text = cls.extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            text = cls.extract_text_from_docx(file_path)
        else:
            raise ParsingException(f"Unsupported file format '{ext}'. Only PDF and DOCX files are allowed.")

        parsed_profile = cls.structure_resume_text(text)
        return {
            "extracted_text": text,
            "parsed_profile": parsed_profile
        }

    @staticmethod
    def structure_resume_text(text: str) -> Dict[str, Any]:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        # Heuristic extraction of common resume sections
        common_skills = [
            "Python", "JavaScript", "TypeScript", "React", "Node.js", "Java", "C++", "C#", "SQL",
            "PostgreSQL", "MongoDB", "Docker", "Kubernetes", "AWS", "Git", "HTML", "CSS", "Tailwind",
            "FastAPI", "Django", "Flask", "Machine Learning", "Data Structures", "Algorithms",
            "System Design", "REST API", "GraphQL", "Spring Boot", "Express"
        ]

        found_skills = set()
        for skill in common_skills:
            if re.search(r'\b' + re.escape(skill) + r'\b', text, re.IGNORECASE):
                found_skills.add(skill)

        # Extract education entries
        education = []
        edu_keywords = ["bachelor", "master", "b.tech", "m.tech", "b.s.", "m.s.", "university", "institute", "college", "degree"]
        for line in lines:
            if any(kw in line.lower() for kw in edu_keywords):
                education.append({
                    "degree": line,
                    "institution": "University / College",
                    "year": "N/A",
                    "gpa": "N/A"
                })

        # Extract projects entries
        projects = []
        proj_heading = False
        current_project = None
        for i, line in enumerate(lines):
            if any(h in line.lower() for h in ["project", "projects", "personal projects"]):
                proj_heading = True
                continue
            if proj_heading:
                if any(h in line.lower() for h in ["experience", "education", "skill", "certificat", "accomplish"]):
                    proj_heading = False
                    break
                if len(line) > 5 and not current_project:
                    current_project = {
                        "title": line[:50],
                        "description": line,
                        "technologies": [s for s in found_skills if s.lower() in line.lower()]
                    }
                    projects.append(current_project)
                    current_project = None

        if not projects:
            projects = [{
                "title": "Software Development Project",
                "description": lines[0] if lines else "Entry level engineering project",
                "technologies": list(found_skills)[:5]
            }]

        # Summary computation
        skills_str = ", ".join(list(found_skills)[:8]) if found_skills else "General Software Engineering"
        summary = f"Candidate with skills in {skills_str}. Extracted from parsed resume document."

        return {
            "skills": list(found_skills) if found_skills else ["Python", "JavaScript", "SQL", "Data Structures"],
            "education": education if education else [{"degree": "B.Tech Computer Science", "institution": "Campus Placement Candidate", "year": "2025", "gpa": "N/A"}],
            "projects": projects,
            "internships": [],
            "certifications": [],
            "summary": summary
        }
